import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.worksheet.page import PageMargins
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
    Paragraph, Spacer, Image as RLImage, HRFlowable)
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from odf.opendocument import OpenDocumentSpreadsheet
from odf.style import Style, TextProperties, TableColumnProperties, TableRowProperties, TableCellProperties, ParagraphProperties
from odf.table import Table as ODFTable, TableColumn, TableRow, TableCell
from odf.text import P
from odf import table as odftable
import base64, io, math, os, tempfile, re
from datetime import date
from PIL import Image as PILImage

LOGO_B64 = "UklGRhoeAABXRUJQVlA4WAoAAAAQAAAA6AIAAwEAQUxQSHAPAAABHARt2yYJf9bffp1BREwAoxQUCsWcBNho54FW6RsDDDl0Ml95ITCNxaytoqETFArIT1i17dS1toRIQAISkFAJSEACEioBCZWABCQgIQ7ycR+nUBI6zldETACk2nbDRh+CIBiCIRhCIRhCIBhCIQhCIRhCIAiCGGQ55zwlmW1EBKTYtsM2hSAIgmAIgmAIghAIghAIghAIgmAIgvAZ3HVS/ff/MiICevuQcqm1yadca835CA7f4CHm2mXCVnP031sunk0mbyX5ryuKpctDuebwPeVSk4fzldwXlEtN1thO/9VEscpK++m/lVxhWW7P7gspVFl0jY9wQU2nlw8zG+rCk+vpPsy7tFBl4Xy6+bKo2ms98xG0qTLzQLM8eZ9eZd6FhSqrv8Jsvqnyd6/5cHtHj85X0bDGuYBTo1/5ymHjdmxURMseWURW6tcruU2jRxZZFI2JwrNiItJPv2U7LldF2XOQoKaaiPTT7Rc9qsSi7+4gqConIi3Sbu2Y6BKd3TigqCfCxe0VPSLfRes8QBT9ROQKW7UDiqL52TkUC4jUsFH0cE5R/msYiglEatinFQwVUT8GhmIDkRp2iRYKNbHgwlCMIFJqkCsSamLDs1OoVhAbMdLioCZWzElRswLsFuIKg5oY0o0BNTPAuwRIi6KKKc/OwLMd2C3AFUQRY+aHgTcEjPhoIWSx58EgWoJVwlsRHGJRZ3BaAmvR0fQ7Nsl1GoLLEtCjW/qbGPU0hJoUZnA09VnMmp2ANymsEtsS78Ww7Gf4i7bgKaHRtDfLCPsZcNqCp4S2pGexLfsZ0GzBviLjEu7YOML+F45tgV2RTeFFzMvXDxCNwVMCo8p2MnmvV84xhODwz9Ja62OsdarY9QNcxuApgU3ZZZp25Rgchrc+lp3HLj+gbgyeEhdVtJMJuebDY+ra389hPOV7BGswA5uiy6h6Hg7PLH3aSawf4LQG77iokp0M5CsFPPx673OYP6BuDXpcU3K+7Uoea7zedgqv7xHMYVdYVMX9lnYGLPX+3IT2PS5r8JSwpuBDPr+Sw3rbN29h9XtiazDDouotH3A5CIu2lTfgKV8jmYM7rKmX/8PlwNJtZT3zezRzWImKqvaQP7kcWL+tLKd/H8zBJ6yptvx2HVDSvuVWv8ZlDu6oqGK7iNRIULTvYp7vnT2sRDW1OunZQduZtYyvUczBiMqK0JCvAxrbr5bra2cPalAMkS5dIhKh9CdLn6/hw/CYp6y19mV8orIi0KUmPzvUtl3J+HqpFGKu/DxaUAx54ZI/o17SUcnl6HcXPYpXVFa0xS5/d6g+snC5+9m/UUkLiqEsdvlv1E121vHyKGl44YrKiqzY5d8d6nudFZ+SraziCoohKlT5MOqno4zpVbJVNaOyoohO+bTDgrOMy63Udg0lKIag0OXjaAL1rFqOpSNLXlLHBq3IOeXzDiP2LKJ5Vj8rttYBMsS4JjdmK6hn0eNa9iugSdUGrUg5WG5kMoN61tBdS14wtQ6QoSTJrScM2bNmO5fzplUbtKKjyL3OEupZQncux7m1fkBeKqjIvRds2bNke9fGp1YFuEVQk5sPY2iU0L1b0lvsBOkSqMnNHeacJcu7Bs2lVQFuBdTk7mQPHRU07/rSL7ETpAsocjsZRF7xcW8Jf8QqwO2/yO0FJj0LqN41YVM7Qbr3LPcfNrEseLtXsFxiFeB2HuX+DqP2AvM/4a52gnTXngecVtHB090r2alWAW7P1GWgN4t+/PL/ZZfcCdIdXzKwwa4WONV9Z5GrALffJCOTYTT4t3sFW+VOkO7V8xBnGS18+/+yTa4C3F6bjGyw7UnT3A/2pXeCdJ9ZhibjdHy6Fzv0KsDt0vEYZxwt2vyf6EfwBGkeqwztsK4FTHPv6BKsAJfDQ8ae5tGg3+4X+ig+QJq/PijYRz94ux8oii3B5S7LWIaBG8yVXC2Q5ox40GUhLfjlXf+eJbicZRkcTWTJruxqgTRXjkc5E2mxSL0kW4LL1SmDO4wc7K20SdYCaY6IRxUrTfadXktwOcoyOlpJga70aoE0N8TDnJkmSn4tf7jcRBnNMLMl2tKr9UOalz7sspMWOoT+qTb74XISZHg2lKGf/Gr8kOajjAuGkpPbt98jS3C6IBlPluokxXXYIy2Q6iGO6zD1STbXaaot3BKcHtq4aquDfLk+p1rCtUDq+U7GZ1sZOV3XXTJynp8miLbSD1yuZeq3cjlIPb5NEIz1Ac2zn2tIb+Q83cmEsHZyeE5z3dLlIPXwOAGby8HmuM7VtDdyHn5NUM31AW+/JHP/aZeD1LNlwstcBg6/ca6tvpHz6DBDNpd+L9jmWurlIPXkvCUHt9x6mXvIb+Q8uc4Q7NXfr0x2y5eD1INlSxTY9upk8qK/kfPcMIUzmGN4LZPtP/1ykHpsngIGny8XZPIZQSPHsdem9Jdrs/UI5KCVU/umKLHLZZbZSwgNZBxKMmMz2caaRy+zP38hyEErZ4YpqsnWixFP9wqigYwz07Z83ouaTF+DkINWjjy3pb9XkemfvygGyDiyboteq8j8rzC0QSsn8r6c70SXPLDGMUDGibIvv1dyTR74+YtDG7Rynt+Y9UaB5Yl3JANknBc25ngfOuWR+y8SbdDKcWljxuuELs98xTJAxnF5Y/rLuCoPtRKLNmjltHNj9CqhyGPHXzADZJxW5+hvT6zyXC7RaIPbJaL7JG55R2F5cv4LZ4D0PNVJsHSfLnk4Uzza4D5s1oT5kM4qC4wIaIB0j1fUvj60yTIbdkQb3B6b5nOOOiaI9mFPBkh3eGvO9juxJ9rgdjgCK2Oicky7MkC6v3dgeUxW7sCuKMDtbwWWDHNiXyZId7cDC2Mu1RptjALc7gjMjamascfOTJDu7ooLY5tmEVujALe7O6w2SBQ/sTkTpHsbYV1jSLGC3VGA29sKK48JejXanwnSnZnkOsUx5lCrEfZHAS5n1KjcmKxVI+zQBGmH8DR3UIyxRalG2CIFuA6p07yDugZVnRphkyZI8/UolhnzIFG5EXZJAS5flJjCGKdSJezTBGlH5HluvTQFxh4aFTxbrQJcvqbeMEMdlBVK2KsJ0k6I8+yQ8qCqDgdslgJcJ4R5uOTmGfwg1qYStusAaQf4iUZAjLFelM1YoF5LcB2AiR65dYIyKOrSPHZMC6Qd0Oehqm0TxEFFlYw1CrYE1wF1opdamZAGsSLVYde0QNr/nxNtsW6CC2O9qNkPLFOx5Q/X/6eJuLSGCeKgpAUnLFSx1g+p/x5mmlrzBDSo6cCZsHlmP5z/bqRpvcYVjHWiYU+EtUrWAq9WrQDpUtu4Y1BSoEYsV7Ml6OU/8pEqwzsGt9Vx8ViwZi3watWLpAkN4/IgJ2u/ItYs2hL06oFOoXmcG5SWdhFWLVoLvFqxoVSd17CCwX1p7F4GS9CLdaJTJw8Lg4KsvbwMWuDVih2lqHQyumJwqnPz4l4GS9CLJztUxmFx1OyLry+DFni12sZaFVlGdazVEpfwMhjptQqUKbKPiovR4uvLIAevVvtlqRKdDO5YrQUu8WVopNd+4I/ENCosR5PvL4McvFqpYJrCa1DFggOX/DI00mt/8KNQBocVTZ7pXZCDaaUT5qXvGFSw5I1LfhkaeK1So63IK2OY1jR4ce+CHEyr1AnzkdfHJCx68+VlaOC1Sg+aW5yXoRWrHrz4d0EOplUabkXbOYTdsuR8fRkaeK1K/Wg+2vqQhHU3XsK7IAfTKifOrczLyIqVO99fhgZeq9IStyrsHMG0tJa4xHdBDqYVynEeYX1EwNoX3+ldaOC1KjvPW9YhAxMWb4lLfhfkYFqhTp5bVRlQsPzFM70LDbxW5SywSxPJ/Y3WZ4HL+S7IwbRCS56nSIr3NYKCkxf3LgzwWoX6FrAktdvYQ8Xgr3dBG0wrbBVMQV7uZg8dJy/hXRjgNQvlFXQ95S720HLz9V3QBqOyldDVEN/EHmoOXo53YYDXLNQuoYlJci97KLr5nQttsD9p1Nilpd/THDQdPD0XA5T4IO0S7FJyyK2NoKvzVlKhDfYnjRrsElJvKdC28YxcDFDig/SrwS4ZXm7kCH2dt5IKbbA/qRVhl4pyQ/NQuCXOzMUAJT5IXoR1DU4+P6Hz4qmp0Ab7kyyLoEsoH7UApS35lYsBSm51WmW8BTj5kDP0XjwtFdrg1H8rypjF/flBcVDcgl+5GA8ZdTzVOfG/ioPuk6enQvsZ+tVhzXeWfxYH9YPfuRgPsayD4dnxHz07GHDyvFKh8xk6KlnVb5GfXA4YcfNWUjEfol2J3V6diPRywI6DZ6RC8ZCWlfApLum8koctN09NxXyIjlqs+zsuiTDnKJipUDxEv1pYlyt/sog4e8h5rlTMp1gWw7t4canLzw6DtoKVCsVDNMqxURy41OTP0yJynpaK+RStctj9cJea/NebxJJ/UqF4inY97Fc51qUm/++w6eLpqZiPsbgB2HAr8meTj0+jWPK7ZELxFPW8A8h1rIViYbnTGUWLZ6RiPkbzJiL9dKs4ziY3V5g1eCuZUDxG6y4i0pJ7GoVcZWC0y+R5p2I+R34fEenn8RgfzyZjOwwbPDUTiufovNPPmgNN5uNZZcJomU/BTMV8kJ2H/WxXDm4CCjFfTSZlsow2T8uE4jmy57zfW805huA+oRBCzmdtMneCaUfBSsV8kC479MNea5cHdxh387RMWCryd5mDpx/W6QU7E1qawLNxKsxb5qNnwlgTeLaNs497gJVEKKsCz5bJMHCZj5EJYlVAzS4NFiaez0oikHUBNauwMxHyfMxMEOsCuowSYWPi+aiJQFYGKCYpsHJ+wCcTxNogGqTBzn0+WiKQ1UFga3QyVHzAygSxOnDNFuxh6T4fPRHI+oCKJdjD1McDdiaI9QEim4E9jF3n45UIZI3gmhHYw9rhAVYSQawRkE3AHvau8zESgawTfNePPQzuH2A1EcQ6gbJ2jWDyMh8zEShKAb6pVgg2dw/gSoRTC0isV4LZywNWIlD0AhWleoDdieejJcIpBviq0UWwfH7AkwgUzYDQteEDtieej54IpxsQ41EXwfr5ATsRKMpBMx7TAzawz8dIhFNPmucjOGEL4wOs5AFFP2ns23MmbGKfj5EIZwGped6ZM2EbjwdQ84BiAsnmedeeCTtZ6vyvH5515o7GOq9fnKuPPkhS+8YNW8SX8sfzVlw8vpk/Hne5Dnw/97XLuUTCl7R9vmddOwO+rUO++nQ1H4QvbQq5VJ6i1xw9vr8pHDnX2m5p9co5BHyhU/gv4YsWVlA4IIQOAACQdQCdASrpAgQBPikUh0KhoQrFLhgMAUJZ27hdT4A/gH4AfoB+eDTkkWg/135gcGFsn+O/wf7Mf0np7OkfAXyAYGWUfPL8a/Sv9b/af7n72P+R/s/yA+dPmAfpt+w/+Q+OHpT/s3oD/Xr9u/dk/4n6Ye6D9oPYA/oP9N//H/H94T/veyj6AX7U///2W/95+5/wO/1T/i/ud7Tf//9gD0AP+314/QD+AfgB+gH8V/f3v8Hi6I+wQVIxP+kFdkWbT510HhMu+6m7H8IkZeeJiuo3k0km/7D3B9qh0vMF18BeeDQkFzPFU/+4PtUOl8YpsAK1+nfXtCchGIzR9+PXjzuqr5+Z5VytQOV6GYEfuN0ByLqwDcQmHnk+Vw2kRR/to+GgcJ6WhSFYcuadUINcMNufuLNA9ILyCw8liz3wUYFy3TtOyoatqRKKRo1QtQN1QfOmbz20SV4WH/f57eoYtVK5l0j3Cng/hEjLwzaky2Iw1sdCKzXgfVJ29ebXM99oEf3jUT47bcMM7aR2aj2hjLjIPoXEXUtrRpwhNcVS6aA2E+3WE2iEmSN8QOzf/+kzVUDxQpGSPpB5LwzKsa4AO6pPo5YpkQyiP5dsjsrFnssBmiDYZJbmg1FyKtaIcPS35g26OdGoPjXq879IPIG1Nym+QIKShSCw8liz2tupPI7hVbQ+P/SC3limLVwwMRGGxhXRgkdpV+XM8Emduuy+6ucgb9P12l84Mei6hGFEkSBWimLVwnIK01h4ZUy88TCwSh79H5rXtVdF75ewmrN6GaApRYZvt/imKYtXCcf62iDBFcKHuTsnSqufO+DUgs9tJRWFuQ9Ph5Pc5wBJBtDeHX0+AEj5Zo87zQ43rRxPhl4ZtODfN7aH29AFDeSfNrV1PrYHrIs7Z6xrfJeCXBAevp3k+kHmL+1G7ao0+VbMOUianDwcSaTQdKOWNIRBX2ffZ9b1AXF1HC/8MkLbuhiJdOcTqoYtWyv6mxB/P1YlGb+iXf20wvLWWIDYL7UhvqZUMoapy2D9Z0Xjv1A3ZJl8zjEmBw54iPp9DjFoofyaK4X/zAOYyYh8YeTix99rsvurnIHZwGIjDY0WKYgXeBXIK/74CSEc/ZzJgpi1cJ0joLD0vpdNAbCfbi5FFTydaCI4j0V4cVdQHOxvuik08CtDQmDbqq6ZgBS9+Pz+srUKrntmulG5LPi1IJPoxhfJsuSZfnrktnH+VLdx3nVKefV8BI+kFyDOe7scgznvPpCAroPEei8vsuJAAP60GakNfeVbKu4/XxN1iFJSMFdEOdkx9YKvf93TObbNxUGF4byN6uK9ir1JV0nGv6+n95a6DGvi/2FaYHWUZXLryDnzZ2djwNHJN8B7CwKO8H+rOn0fdxtjrDVCix5z1CFY9nFFfBQXzCuk0CUE+FfgbfvKg4N0F5/QqOoGa4Un9z56ak/Cmi3Z4jfdsMBKXzoglbApkS4wU4ScY69Rk0AUMikz95UFfEfV4cGLbXTbUK+Sjava3IhAFQ1kHz///oV3nGxQSqX2FKuSdvDPHvfdG3bBXn3lO7/66JrSDSBHtOhPKjYrugwoY/y5bGB0Jr/2PWMLtyH5lBv4ABego4o/syTLl66vlzUFMdUh6+YJVt0ANHHcOP4BQfr8YZgeG5cDywuP7mWRANHeHFijabe0XWVWIUSXUmdFR4rpYs/R0nIJHqJO4Ye8NQBsidULpFgacrJ99SAF0Z/BDIZYx+/UsQeX2MfP82MaeFaBxtoPv7ydSE1KJ80H9gg0uQrvf6f221L+VOGJbjXR/PguvJIS2Bxi+CMaoiNlbhFvGfvCj96sDvHXmFogNZjtA4b2SseHdMJ80AK+fuYtUIbJzIoYHn8usa++VnelOd6nIIbhJe7f3+ijMTxQM/MYxf16Cd5G/gmdtqiXtBKZGLRamZDDbXJOHn4w5LAeqxu9dBuqW5w41L9ImBqo2wtDfBZvlDaB+XcqYLT3iT4hHkYyBF4mS+vjGXp+icoDn+EJRv5brgGIF595R39xxeY/eTo48i1jkDcljMy1YxbYTBC4oU/3hNq+eKtBvWBZKSJYpSzm+36gvdKEh8/eDC0EO2WprMCpeAPbOG21str9UYkDFyJ/oecd5GtQkaJyOjQlhCXdxF36wQCQ38wqsVXUG49MoWa8zGxeCwuKyECkvIM8X2slN+/CCzwRZOwHuaEtq47e6GIbCE/DvUVLveFCISedkFbnwFFytVMaGEoBWGPd00VQUvUnADPzXrJyfIkZDNhbrzg9PAQm5tf9q0iYlL8wXKainA3T4HZiVj+EqK4mxLxZ2ggTgG+gT1DL94MtyC6EfN7DmGHMMJfpgy/H+VZavkEGT1q6T+HXoCWsFCUDELsh+/UsQe/8wvic9Sc08V17x5F3Sp2iB6A8YBgW2mwTtBE5ZIi+dmbn1CPq1uyUt6fwzktG+x6sggkrsz3fMmd4/fGMvT9E5QHP8ISjfy3XAMR7/eXsb3EmF0kwrrrEpN4aYdl58RBGUZzF4ITa/QgJjryJjBEmApoAADi5ele2up8/oUykBOVO/OvronOqgbqv3Pim6z6KdpwuioXIy/JFGe6zJbnAu9S4oK/iLfnCIekWR6bQKPxLWQ4NOZKy3994K6PNVqJI1oeLpV2xXxXe0sUd1HAz4K8uULOKLFyZaaN8zYM47M3BNdc8JPzfzSTH4WUY4Wqq4G6vVglV2SUlApPvDBY7Jr3jh0/Xluv2zfdebpH4n4ECfl1onUA8nxOQbGkqZYe2E8n3laS/e129xmtXC7+rGJEfd8rPNln1VtQNlJuQE6l3GHPdPpcG6y90CY9W7yxZ1tjEFFQuRl+Lp1zcqsVXUG3GebM0U1OYQLMfyyS+db85rPxIay74s9riPg+l4fheDJ6rnFvy78J6AdSC8mM++8pEXOJ39+buAXwjZw0ZSyqSTQ56W+V0NMRqP7pW1/B4DPRfMQt8TGklAKHg1kgLxrHNyoIImKjTmAxrSNxTZp2hIeSjbBfJHs2xKDP4Y7Lx8ybhT8fPlHVJMMw2L6p4Itz+FZg0MGIbd0DQ0AHf3k62XKuDvOWGmI+tojk7L4p79xY7C31cdvc7oP73AA2doqW6W+eCkE13+kZbp/uwx3QPgtVDhjQO1dCxalY5wVU75Wyefz9Wh8jMnIudBfxWETNwZFqR5C7PniljT2O9IRlKiQLr2bcyLitX6VsYXT7F9BpgsAVsmO8s3tpGXD7sFesR45BmvZHLkqno2/uOxCdkLaQKQU0Du2hoGoA/Ao5CbjVFP9hO75CALwAzq1f6KaDYf0eqCs4RJS/eHT7wgXLbZQhyailUR0Py+EeixRXP6WKLzPi+OTVthjaEMNeEqz9Co0+m856pGvcwAUAtqDqXHh+89gNJ3nwC2Xq/wnCr5AlpBgqEfC6PEskjL3DI65oa+Sz9OSj74C3RvpREDubX66xKThmXR/fci/bFnY/PG+8LqeQKIWz0utfA57NuG3MaD7ycylSxQ3/xhh1kNWeebPi2DfxXS3zwUg1f/X2brnIU43pmP7xWfdHUHjR+cCc9ShfP3WTfTHQS5QpzAY1S8acXbPEuHQH34rCJm4Mi1I8hdnzxSxp7HekM/iIdTbpj4fv5lWBM7/uAV8gcSpuoif5vbqPBu79aHkiCWlyPeGP5VEmK+TGlZGaD95R39xxeF222B/mdO/bn84MXn3fDE/16SmkHkLJzGX3G8N/eF5EGb9NLaA/xuuhoy8u81nc9cr9dwdnuVa670bGLZWCstAeuoS2bnX/SES6ZWcntP9OZDbYbs5kuwH8q7oiTMlbSaVAjn/0rCJm4Mi1QL7oLnmYMVYV+a7pZehU+919uqiUEWs3706WrOCXVkimzTxExuUp3dcPczj4aR5fKBMLPrSEWDCesKId2lyQFwc1Q5XP6bBKtubkRpXf7ypXfyR56v5/ccXws8kePPzgYqgE9USG0uu0ignSDILhcgxO6wP7w6feFWexyUkSzo4h9REJA8TaoeUE6C5cf0bSmhOvizqYcLz0a7PIEGkYlm8tOqLGEpfEhwbwzsliRIRYBx7dEPPNnxZ8DKPFX8N86slezaM3YSYuSOUAo/FbVPRUk4N2epNors+FUClj0ES6YKT+zDKFQPMRDGfcHA+JWSBTr+A+Mkbthqigh+Beu6D/+453ENLfE8jIikICKYp1qcpYs/eR0IL15C8r8g6e+Be42MJo5hEvpt5/ITkTsdWh5dYpfeDbgZQKXsmVLy5/T3LufBun9Kh7zk6gyhCPqH5+MOSytA9cOhH/iJrL1k/Qwi7Rjb7gIwV1UvaAR9UZz+lih29NGIYz6yAHL94FjDQXf7EFSam7eYQVB7vnZHQUGL36UyAzpgWpcl4ABRKAAw2Q/9/g5Gdv1bgK3I11mPof3gxPiAakj+nDxt91ELxkC/5pAHei6B+Kz3skUfpTaFwqU9SeMBCWfp0gffAW+NbNZ0rFqVjnBWKSAnVtXFA7zFE+8XdaGV8JaY3aAismX/EWzQlhEm07Sjart7hRb/7/L2mrdf9ST1h1q2p/QoMLMXAc1S37BFCHB+kzCZEG/tCNNVVR9Np+AhrybNm/sgpyLFI94MiYbfvDp94Xi7XEVJLiXNsq/NaI7rZG5VKSLFsy8pHttRTdx+jKrYL4vyZ2EQHxZki06eGNnT9yZnyxDBcHSObol366vltKY6MrUn8eMs7YguSs2oS4BygpKs54EGMBddUTSfq/rpPtLzNvs9cPqKP4msmoJxh0ssTljaD/VOfG3RCeSFgFGEU0X0q3Sd8hmIXrcrX4TXxO/lnfT5FC+99MrtRC37wjaPX3Pye5CtRTRfSsT7Xg8F/7X/eswPKKMRRkZiVkqz2qFkpcMA/P4ixhG7ag+LtPCEgZpGBRcwHj+JFkcMw52Qo/vBe4hXWINsBgFVVH1gIs7YAAAAAAAAAAAAA=="

# ── Farby ─────────────────────────────────────────────────────────────────────
PRI        = "3D2B6B"
WHITE      = "FFFFFF"
PARAM_ODD  = "F7F7F7"
PARAM_EVEN = "FFFFFF"
BORDER     = "CCCCCC"
TEXT       = "1A1A1A"
MUTED      = "777777"
META_LBL   = "F5F5F5"
META_VAL   = "FFFFFF"
SECTION_BG = "E8E0F0"
PASS_COL   = "2E7D32"
FAIL_COL   = "C62828"
INFO_COL   = "1565C0"

PRI_RL    = colors.HexColor("#3D2B6B")
LIGHT_RL  = colors.HexColor("#EDE7F6")
MUTED_RL  = colors.HexColor("#777777")
TEXT_RL   = colors.HexColor("#1A1A1A")
WHITE_RL  = colors.white
BORDER_RL = colors.HexColor("#CCCCCC")
ODD_RL    = colors.HexColor("#F0EDF8")

def _bdr(c=BORDER):
    s = Side(style="thin", color=c)
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(c):
    return PatternFill(fill_type="solid", fgColor=c)

def _est_h(txt, w, base=15, mn=15):
    if not txt: return mn
    ch = max(1, int(w * 1.15))
    lines = max(1, (len(str(txt)) + ch - 1) // ch)
    return max(mn, base * lines)

def _logo_png():
    logo_bytes = base64.b64decode(LOGO_B64)
    pil = PILImage.open(io.BytesIO(logo_bytes)).convert("RGBA")
    bg = PILImage.new("RGBA", pil.size, (255,255,255,255))
    bg.paste(pil, mask=pil.split()[3])
    final = bg.convert("RGB")
    tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    final.save(tmp.name, "PNG", optimize=False)
    return tmp.name

def _calc_layout(data):
    layout = data.get("layout", {})
    params = data.get("parameters", [])
    has_unit   = layout.get("has_unit_column", False)
    has_method = layout.get("has_method_column", False)
    has_sect   = layout.get("has_sections", False)
    mx_name   = max((len(str(p.get("name","")))    for p in params), default=20)
    mx_mm     = max((len(str(p.get("min_max",""))) for p in params), default=15)
    mx_res    = max((len(str(p.get("result","")))  for p in params), default=12)
    mx_unit   = max((len(str(p.get("unit","")))    for p in params), default=0) if has_unit else 0
    mx_method = max((len(str(p.get("method","")))  for p in params), default=0) if has_method else 0
    mx_name   = max(mx_name,   layout.get("max_param_chars",   0))
    mx_mm     = max(mx_mm,     layout.get("max_minmax_chars",  0))
    mx_res    = max(mx_res,    layout.get("max_result_chars",  0))
    return {"has_unit": has_unit, "has_method": has_method, "has_sections": has_sect,
            "mx_name": mx_name, "mx_mm": mx_mm, "mx_res": mx_res,
            "mx_unit": mx_unit, "mx_method": mx_method}

def _meta_pairs(data):
    std_l = [
        ("Common Name",       data.get("commonName","")),
        ("Batch Number",      data.get("batchNumber","")),
        ("Species",           data.get("species","")),
        ("Supplier",          data.get("supplier","")),
        ("Country of Origin", data.get("countryOfOrigin","")),
    ]
    std_r = [
        ("Manufacture Date",  data.get("manufacturingDate","")),
        ("Retest / Expiry",   data.get("retestDate","")),
        ("Product Code",      data.get("productCode","")),
        ("Shelf Life",        data.get("shelfLife","")),
        ("Allergens",         data.get("allergens","")),
    ]
    extra = data.get("extraMeta", {}) or {}
    ex = list(extra.items())
    for k in range(0, len(ex), 2):
        std_l.append(ex[k])
        std_r.append(ex[k+1] if k+1 < len(ex) else ("",""))
    while len(std_l) < len(std_r): std_l.append(("",""))
    while len(std_r) < len(std_l): std_r.append(("",""))
    return list(zip(std_l, std_r))

# ═══════════════════════════════════════════════════════════════════════════════
# VERIFIKÁCIA
# ═══════════════════════════════════════════════════════════════════════════════

def _parse_num(s):
    if not s: return None
    m = re.search(r"[<>]?\s*([\d]+\.?[\d]*)", str(s).replace(",","."))
    return float(m.group(1)) if m else None

def verify_parameters(parameters):
    result = []
    issues = []
    for p in parameters:
        p = dict(p)
        name    = p.get("name", "")
        min_max = str(p.get("min_max", ""))
        res_str = str(p.get("result", ""))
        status  = p.get("status", "pass")
        note = ""
        ver  = "ok"

        if status == "info" or res_str.lower() in ["compliant","not tested","n/t","n.t.",""]:
            p["verification_status"] = "info"
            p["verification_note"]   = ""
            result.append(p)
            continue

        if re.search(r"not\s+detected|^nd$|^negative$", res_str.lower()):
            p["verification_status"] = "ok"
            p["verification_note"]   = ""
            result.append(p)
            continue

        res_num = _parse_num(res_str)
        min_val = max_val = None

        if "/" in min_max:
            parts = min_max.split("/")
            left  = parts[0].strip()
            right = parts[1].strip() if len(parts) > 1 else ""
            if left not in ["-", ""]:
                min_val = _parse_num(left)
            if right.lower() not in ["nd","not detected","-",""]:
                max_val = _parse_num(right)
        elif re.match(r"[>≥]=?", min_max):
            min_val = _parse_num(min_max)
        elif re.match(r"[<≤]=?", min_max):
            max_val = _parse_num(min_max)

        if res_num is not None:
            if min_val is not None and res_num < min_val:
                ver  = "error"
                note = f"Hodnota {res_num} je pod minimom {min_val}"
                issues.append({"param": name, "note": note})
                p["status"] = "fail"
            elif max_val is not None and res_num > max_val:
                ver  = "error"
                note = f"Hodnota {res_num} presahuje maximum {max_val}"
                issues.append({"param": name, "note": note})
                p["status"] = "fail"
            else:
                ver = "ok"
        else:
            ver = "info"

        if not res_str.strip():
            ver  = "warning"
            note = "Chýba hodnota výsledku"
            issues.append({"param": name, "note": note})

        p["verification_status"] = ver
        p["verification_note"]   = note
        result.append(p)
    return result, issues

def verify_meta(data):
    issues = []
    for field in ["commonName","batchNumber","manufacturingDate","retestDate"]:
        if not data.get(field,"").strip():
            issues.append({"field": field, "note": f"Chýba povinné pole: {field}"})
    mfg = data.get("manufacturingDate","")
    exp = data.get("retestDate","")
    if mfg and exp:
        my = re.search(r"(20\d{2})", mfg)
        ey = re.search(r"(20\d{2})", exp)
        if my and ey and int(ey.group(1)) < int(my.group(1)):
            issues.append({"field":"dates","note":"Dátum expirácie je pred dátumom výroby"})
    return issues

# ═══════════════════════════════════════════════════════════════════════════════
# XLSX
# ═══════════════════════════════════════════════════════════════════════════════

def generate_xlsx(data, output_path):
    n      = Side(style=None)
    thick  = Side(style="medium", color=PRI)
    thin_p = Side(style="thin",   color="D5C5EC")
    lo     = _calc_layout(data)
    wA, wB, wC, wD, wE = 20, 24, 20, 18, 8

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Certificate of Analysis"
    for col, w in zip("ABCDE", [wA,wB,wC,wD,wE]):
        ws.column_dimensions[col].width = w

    # Header
    ws.row_dimensions[1].height = 65
    ws.merge_cells("A1:B1"); ws.merge_cells("C1:E1")
    lp = _logo_png()
    xl = XLImage(lp); xl.width=130; xl.height=46; xl.anchor="A1"
    ws.add_image(xl)
    ws["C1"].value = ("ORIN Group s.r.o.\nNa Štepnici 1379/1, 960 01 Zvolen, Slovakia\n"
                      "ID: 46294490  |  VAT: SK2023327735\ninfo@privatebrand.eu")
    ws["C1"].font      = Font(name="Calibri", size=8, color=MUTED)
    ws["C1"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    for col in "ABCDE":
        ws[f"{col}1"].border = Border(bottom=thick, left=n, right=n, top=n)

    # Title
    ws.row_dimensions[2].height = 24
    ws.merge_cells("A2:E2")
    ws["A2"].value     = "CERTIFICATE OF ANALYSIS"
    ws["A2"].font      = Font(name="Calibri", size=12, bold=True, color=PRI)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Product name
    ws.row_dimensions[3].height = 30
    ws.merge_cells("A3:E3")
    ws["A3"].value     = data.get("commonName","")
    ws["A3"].font      = Font(name="Calibri", size=15, bold=True, color=PRI)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for col in "ABCDE":
        ws[f"{col}3"].border = Border(bottom=thin_p, left=n, right=n, top=n)

    # Meta
    row = 4
    for (l1,v1),(l2,v2) in _meta_pairs(data):
        h = max(_est_h(v1,wB), _est_h(v2,wD+wE), 15)
        ws.row_dimensions[row].height = h
        for col,val,lbl in [("A",l1,True),("B",v1,False),("C",l2,True)]:
            c = ws[f"{col}{row}"]
            c.value = val
            c.font  = Font(name="Calibri", size=8.5, bold=lbl, color=MUTED if lbl else TEXT)
            c.fill  = _fill(META_LBL if lbl else META_VAL)
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            c.border = _bdr()
        ws.merge_cells(f"D{row}:E{row}")
        c = ws[f"D{row}"]
        c.value = v2
        c.font  = Font(name="Calibri", size=8.5, color=TEXT)
        c.fill  = _fill(META_VAL)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        c.border = _bdr()
        row += 1

    # Description
    desc = data.get("description","")
    if desc:
        ws.row_dimensions[row].height = max(_est_h(desc,90),20)
        ws.merge_cells(f"A{row}:E{row}")
        c = ws[f"A{row}"]
        c.value = f"Description: {desc}"
        c.font  = Font(name="Calibri", size=8, italic=True, color=MUTED)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        c.border = _bdr()
        row += 1

    # Spacer
    ws.row_dimensions[row].height = 8
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"].fill = _fill(WHITE)
    row += 1

    # Param header
    ws.row_dimensions[row].height = 20
    ws.merge_cells(f"A{row}:B{row}")
    if lo["has_unit"]:
        hdrs = [("A","PARAMETER","left"),("C","UNIT","center"),
                ("D","MIN / MAX","left"),("E","RESULT","center")]
    elif lo["has_method"]:
        hdrs = [("A","PARAMETER","left"),("C","SPECIFICATION","left"),
                ("D","RESULT","center"),("E","METHOD","center")]
    else:
        hdrs = [("A","PARAMETER","left"),("C","MIN / MAX","left"),
                ("D","RESULT","center"),("E","","center")]
    for col,label,align in hdrs:
        c = ws[f"{col}{row}"]
        c.value = label
        c.font  = Font(name="Calibri", size=8, bold=True, color=WHITE)
        c.fill  = _fill(PRI)
        c.alignment = Alignment(horizontal=align, vertical="center",
                                indent=1 if align=="left" else 0)
        c.border = _bdr(PRI)
    row += 1

    # Param rows
    current_section = None
    for i, p in enumerate(data.get("parameters",[])):
        if lo["has_sections"] and p.get("section") and p["section"] != current_section:
            current_section = p["section"]
            ws.row_dimensions[row].height = 16
            ws.merge_cells(f"A{row}:E{row}")
            c = ws[f"A{row}"]
            c.value = current_section.upper()
            c.font  = Font(name="Calibri", size=8, bold=True, color=PRI)
            c.fill  = _fill(SECTION_BG)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=2)
            c.border = _bdr("B0A0C8")
            row += 1

        ws.row_dimensions[row].height = max(_est_h(p.get("name",""), wA+wB), 15)
        bg = PARAM_ODD if i % 2 == 0 else PARAM_EVEN
        st   = p.get("status","pass")
        icon = "✓" if st=="pass" else "✗" if st=="fail" else "·"
        icol = PASS_COL if st=="pass" else FAIL_COL if st=="fail" else INFO_COL

        ws.merge_cells(f"A{row}:B{row}")
        c = ws[f"A{row}"]
        c.value = p.get("name","")
        c.font  = Font(name="Calibri", size=9, bold=True, color=TEXT)
        c.fill  = _fill(bg)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        c.border = _bdr()

        if lo["has_unit"]:
            vals = [("C",p.get("unit",""),False,MUTED,"center"),
                    ("D",p.get("min_max",""),False,MUTED,"left"),
                    ("E",p.get("result",""),True,TEXT,"center")]
        elif lo["has_method"]:
            vals = [("C",p.get("min_max",""),False,MUTED,"left"),
                    ("D",p.get("result",""),True,TEXT,"center"),
                    ("E",p.get("method",""),False,MUTED,"center")]
        else:
            vals = [("C",p.get("min_max",""),False,MUTED,"left"),
                    ("D",p.get("result",""),True,TEXT,"center"),
                    ("E",icon,True,icol,"center")]

        for col,val,bold,fg,align in vals:
            c = ws[f"{col}{row}"]
            c.value = val
            c.font  = Font(name="Calibri", size=9, bold=bold, color=fg)
            c.fill  = _fill(bg)
            c.alignment = Alignment(horizontal=align, vertical="center",
                                    wrap_text=True, indent=1 if align=="left" else 0)
            c.border = _bdr()
        row += 1

    # Notes + storage
    for txt in [data.get("notes",""),
                ("Storage: " + data.get("storageConditions","")).strip()]:
        if txt and txt not in ["Storage: ",""]:
            ws.row_dimensions[row].height = max(_est_h(txt,90),18)
            ws.merge_cells(f"A{row}:E{row}")
            c = ws[f"A{row}"]
            c.value = txt
            c.font  = Font(name="Calibri", size=8, italic=True, color=MUTED)
            c.alignment = Alignment(horizontal="left", vertical="center",
                                    wrap_text=True, indent=1)
            c.border = _bdr()
            row += 1

    # Footer
    row += 1
    ws.row_dimensions[row].height = 46
    ws.merge_cells(f"A{row}:B{row}")
    c = ws[f"A{row}"]
    c.value = "QA Manager\n\n________________________\nSignature & Date"
    c.font  = Font(name="Calibri", size=9, color=TEXT)
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    ws.merge_cells(f"C{row}:E{row}")
    c = ws[f"C{row}"]
    c.value = (f"ORIN Group s.r.o. — Internal QC Document\n"
               f"Generated: {date.today().strftime('%d.%m.%Y')}")
    c.font  = Font(name="Calibri", size=8, color=MUTED)
    c.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True, indent=1)
    thick_f = Side(style="medium", color=PRI)
    for col in "ABCDE":
        ws[f"{col}{row}"].border = Border(top=thick_f, left=n, right=n, bottom=n)

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage  = True
    ws.page_setup.fitToWidth = 1
    ws.page_margins = PageMargins(left=0.55, right=0.55, top=0.55, bottom=0.55)
    wb.save(output_path)
    try:
        os.unlink(lp)
    except:
        pass
    return output_path

# ═══════════════════════════════════════════════════════════════════════════════
# ODS — priamo cez odfpy, bez LibreOffice
# ═══════════════════════════════════════════════════════════════════════════════

def generate_ods(data, output_path):
    doc = OpenDocumentSpreadsheet()
    lo  = _calc_layout(data)

    def ods_style(name, bold=False, size="9pt", color="#1A1A1A",
                  bg=None, align="left", italic=False, border=True):
        st = Style(name=name, family="table-cell")
        cp = {"wrapoption":"wrap", "verticalalign":"middle",
              "paddingleft":"0.08cm", "paddingright":"0.08cm",
              "paddingtop":"0.04cm", "paddingbottom":"0.04cm"}
        if bg: cp["backgroundcolor"] = f"#{bg}"
        if border: cp["border"] = "0.5pt solid #CCCCCC"
        st.addElement(TableCellProperties(**cp))
        fp = {"fontsize": size, "color": color}
        if bold:   fp["fontweight"] = "bold"
        if italic: fp["fontstyle"]  = "italic"
        st.addElement(TextProperties(**fp))
        st.addElement(ParagraphProperties(textalign=align))
        doc.styles.addElement(st)
        return name

    def rs(name, h_cm):
        s = Style(name=name, family="table-row")
        s.addElement(TableRowProperties(rowheight=f"{h_cm}cm",
                                         useoptimalrowheight="false"))
        doc.styles.addElement(s)
        return name

    def cs(name, w_cm):
        s = Style(name=name, family="table-column")
        s.addElement(TableColumnProperties(columnwidth=f"{w_cm}cm"))
        doc.styles.addElement(s)
        return name

    # Styly
    S = {
        "hdr_co": ods_style("hdr_co", size="8pt", color="#777777", align="right", border=False),
        "title":  ods_style("title",  bold=True, size="12pt", color="#3D2B6B", align="center", border=False),
        "prod":   ods_style("prod",   bold=True, size="15pt", color="#3D2B6B", align="center", border=False),
        "ml":     ods_style("ml",     bold=True, size="8.5pt", color="#777777", bg="F5F5F5"),
        "mv":     ods_style("mv",     size="8.5pt", color="#1A1A1A", bg="FFFFFF"),
        "th":     ods_style("th",     bold=True, size="8pt", color="#FFFFFF", bg=PRI),
        "th_c":   ods_style("th_c",   bold=True, size="8pt", color="#FFFFFF", bg=PRI, align="center"),
        "p_odd":  ods_style("p_odd",  bold=True, size="9pt", color="#1A1A1A", bg="F0EDF8"),
        "p_even": ods_style("p_even", bold=True, size="9pt", color="#1A1A1A", bg="FFFFFF"),
        "mm_odd": ods_style("mm_odd", size="9pt", color="#777777", bg="F0EDF8"),
        "mm_eve": ods_style("mm_eve", size="9pt", color="#777777", bg="FFFFFF"),
        "r_odd":  ods_style("r_odd",  bold=True, size="9pt", color="#1A1A1A", bg="F0EDF8", align="center"),
        "r_eve":  ods_style("r_eve",  bold=True, size="9pt", color="#1A1A1A", bg="FFFFFF", align="center"),
        "ic_p":   ods_style("ic_p",   bold=True, size="9pt", color="#2E7D32", bg="F0EDF8", align="center"),
        "ic_pe":  ods_style("ic_pe",  bold=True, size="9pt", color="#2E7D32", bg="FFFFFF", align="center"),
        "ic_f":   ods_style("ic_f",   bold=True, size="9pt", color="#C62828", bg="F0EDF8", align="center"),
        "ic_fe":  ods_style("ic_fe",  bold=True, size="9pt", color="#C62828", bg="FFFFFF", align="center"),
        "sect":   ods_style("sect",   bold=True, size="8pt",  color="#3D2B6B", bg="E8E0F0"),
        "notes":  ods_style("notes",  italic=True, size="8pt", color="#777777"),
        "qa":     ods_style("qa",     size="9pt", color="#1A1A1A", border=False),
        "gen":    ods_style("gen",    size="8pt", color="#777777", align="right", border=False),
        "empty":  ods_style("empty",  border=False),
    }

    # Row heights
    RH = {
        "hdr":  rs("rh_hdr",  1.8),
        "titl": rs("rh_tit",  0.8),
        "prod": rs("rh_pro",  1.0),
        "meta": rs("rh_met",  0.55),
        "tall": rs("rh_tal",  1.0),
        "spc":  rs("rh_spc",  0.3),
        "th":   rs("rh_th",   0.65),
        "par":  rs("rh_par",  0.55),
        "not":  rs("rh_not",  0.7),
        "foot": rs("rh_foo",  1.5),
    }

    # Col widths (5 cols)
    CW = [cs(f"cw{i}", w) for i,w in enumerate([4.4, 5.3, 4.4, 4.0, 1.8])]

    def cell(val, sty, span=1):
        tc = TableCell(stylename=sty)
        if span > 1:
            tc.setAttribute("numbercolumnsspanned", str(span))
        if val:
            tc.addElement(P(text=str(val)))
        return tc

    def cov(n=1):
        return [odftable.CoveredTableCell() for _ in range(n)]

    sht = ODFTable(name="Certificate of Analysis")
    for cw in CW:
        sht.addElement(TableColumn(stylename=cw))

    # Row 1: Header
    r = TableRow(stylename=RH["hdr"])
    r.addElement(cell("", S["empty"], span=2)); [r.addElement(c) for c in cov(1)]
    r.addElement(cell("ORIN Group s.r.o.  |  Na Štepnici 1379/1, 960 01 Zvolen  |  "
                      "ID: 46294490  |  VAT: SK2023327735  |  info@privatebrand.eu",
                      S["hdr_co"], span=3))
    [r.addElement(c) for c in cov(2)]
    sht.addElement(r)

    # Row 2: Title
    r = TableRow(stylename=RH["titl"])
    r.addElement(cell("CERTIFICATE OF ANALYSIS", S["title"], span=5))
    [r.addElement(c) for c in cov(4)]
    sht.addElement(r)

    # Row 3: Product
    r = TableRow(stylename=RH["prod"])
    r.addElement(cell(data.get("commonName",""), S["prod"], span=5))
    [r.addElement(c) for c in cov(4)]
    sht.addElement(r)

    # Meta rows
    for (l1,v1),(l2,v2) in _meta_pairs(data):
        h = RH["tall"] if len(str(v1))>25 or len(str(v2))>20 else RH["meta"]
        r = TableRow(stylename=h)
        r.addElement(cell(l1, S["ml"]))
        r.addElement(cell(v1, S["mv"]))
        r.addElement(cell(l2, S["ml"]))
        r.addElement(cell(v2, S["mv"], span=2))
        r.addElement(odftable.CoveredTableCell())
        sht.addElement(r)

    # Spacer
    r = TableRow(stylename=RH["spc"])
    r.addElement(cell("", S["empty"], span=5))
    [r.addElement(c) for c in cov(4)]
    sht.addElement(r)

    # Param header
    r = TableRow(stylename=RH["th"])
    if lo["has_unit"]:
        r.addElement(cell("PARAMETER", S["th"], span=2)); r.addElement(odftable.CoveredTableCell())
        r.addElement(cell("UNIT",      S["th_c"]))
        r.addElement(cell("MIN / MAX", S["th"]))
        r.addElement(cell("RESULT",    S["th_c"]))
    elif lo["has_method"]:
        r.addElement(cell("PARAMETER",    S["th"], span=2)); r.addElement(odftable.CoveredTableCell())
        r.addElement(cell("SPECIFICATION",S["th"]))
        r.addElement(cell("RESULT",       S["th_c"]))
        r.addElement(cell("METHOD",       S["th_c"]))
    else:
        r.addElement(cell("PARAMETER", S["th"], span=2)); r.addElement(odftable.CoveredTableCell())
        r.addElement(cell("MIN / MAX", S["th"]))
        r.addElement(cell("RESULT",    S["th_c"]))
        r.addElement(cell("",          S["th_c"]))
    sht.addElement(r)

    # Param rows
    current_section = None
    for i, p in enumerate(data.get("parameters",[])):
        if lo["has_sections"] and p.get("section") and p["section"] != current_section:
            current_section = p["section"]
            r = TableRow(stylename=RH["par"])
            r.addElement(cell(current_section.upper(), S["sect"], span=5))
            [r.addElement(c) for c in cov(4)]
            sht.addElement(r)

        odd  = i % 2 == 0
        st   = p.get("status","pass")
        icon = "✓" if st=="pass" else "✗" if st=="fail" else "·"
        ps   = S["p_odd"]  if odd else S["p_even"]
        ms   = S["mm_odd"] if odd else S["mm_eve"]
        rs2  = S["r_odd"]  if odd else S["r_eve"]
        ics  = (S["ic_p"] if odd else S["ic_pe"]) if st=="pass" else (S["ic_f"] if odd else S["ic_fe"])

        r = TableRow(stylename=RH["par"])
        r.addElement(cell(p.get("name",""), ps, span=2))
        r.addElement(odftable.CoveredTableCell())
        if lo["has_unit"]:
            r.addElement(cell(p.get("unit",""),   ms))
            r.addElement(cell(p.get("min_max",""),ms))
            r.addElement(cell(p.get("result",""), rs2))
        elif lo["has_method"]:
            r.addElement(cell(p.get("min_max",""),ms))
            r.addElement(cell(p.get("result",""), rs2))
            r.addElement(cell(p.get("method",""), ms))
        else:
            r.addElement(cell(p.get("min_max",""),ms))
            r.addElement(cell(p.get("result",""), rs2))
            r.addElement(cell(icon, ics))
        sht.addElement(r)

    # Notes + storage
    for txt in [data.get("notes",""),
                ("Storage: " + data.get("storageConditions","")).strip()]:
        if txt and txt not in ["Storage: ",""]:
            r = TableRow(stylename=RH["not"])
            r.addElement(cell(txt, S["notes"], span=5))
            [r.addElement(c) for c in cov(4)]
            sht.addElement(r)

    # Footer
    r = TableRow(stylename=RH["foot"])
    r.addElement(cell("QA Manager\n\n________________________\nSignature & Date",
                      S["qa"], span=2))
    r.addElement(odftable.CoveredTableCell())
    r.addElement(cell(f"ORIN Group s.r.o. — Internal QC Document\n"
                      f"Generated: {date.today().strftime('%d.%m.%Y')}",
                      S["gen"], span=3))
    [r.addElement(c) for c in cov(2)]
    sht.addElement(r)

    doc.spreadsheet.addElement(sht)
    doc.save(output_path)
    return output_path

# ═══════════════════════════════════════════════════════════════════════════════
# PDF
# ═══════════════════════════════════════════════════════════════════════════════

def generate_pdf(data, output_path):
    W, H   = A4
    MARGIN = 16*mm

    def ps(name, **kw):
        return ParagraphStyle(name, **kw)

    ST = {
        "co":  ps("co",  fontSize=7.5, textColor=MUTED_RL, leading=12, alignment=TA_RIGHT),
        "ti":  ps("ti",  fontSize=13,  textColor=PRI_RL, fontName="Helvetica-Bold",
                  alignment=TA_CENTER, spaceAfter=3),
        "pr":  ps("pr",  fontSize=16,  textColor=PRI_RL, fontName="Helvetica-Bold",
                  alignment=TA_CENTER, spaceAfter=10),
        "ml":  ps("ml",  fontSize=8,   textColor=MUTED_RL, fontName="Helvetica-Bold"),
        "mv":  ps("mv",  fontSize=8.5, textColor=TEXT_RL),
        "th":  ps("th",  fontSize=8,   textColor=WHITE_RL, fontName="Helvetica-Bold"),
        "thc": ps("thc", fontSize=8,   textColor=WHITE_RL, fontName="Helvetica-Bold",
                  alignment=TA_CENTER),
        "pn":  ps("pn",  fontSize=8.5, textColor=TEXT_RL, fontName="Helvetica-Bold", leading=11),
        "pv":  ps("pv",  fontSize=8.5, textColor=MUTED_RL, leading=11),
        "pc":  ps("pc",  fontSize=8.5, textColor=TEXT_RL,  fontName="Helvetica-Bold",
                  leading=11, alignment=TA_CENTER),
        "no":  ps("no",  fontSize=8,   textColor=MUTED_RL, fontName="Helvetica-Oblique", leading=11),
        "qa":  ps("qa",  fontSize=9,   textColor=TEXT_RL),
        "ge":  ps("ge",  fontSize=7.5, textColor=MUTED_RL, alignment=TA_RIGHT),
        "se":  ps("se",  fontSize=8,   textColor=PRI_RL, fontName="Helvetica-Bold"),
    }

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
        leftMargin=MARGIN, rightMargin=MARGIN,
        topMargin=MARGIN, bottomMargin=MARGIN)
    cw = W - 2*MARGIN
    story = []

    # Header
    logo_b = base64.b64decode(LOGO_B64)
    logo_i = RLImage(io.BytesIO(logo_b), width=55*mm, height=18*mm)
    logo_i.hAlign = "LEFT"
    comp = Paragraph(
        "<b>ORIN Group s.r.o.</b><br/>"
        "Na Štepnici 1379/1, 960 01 Zvolen, Slovakia<br/>"
        "ID: 46294490  |  VAT: SK2023327735<br/>"
        "info@privatebrand.eu", ST["co"])
    hdr = Table([[logo_i, comp]], colWidths=[cw*0.45, cw*0.55])
    hdr.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("ALIGN",(1,0),(1,0),"RIGHT"),
        ("BOTTOMPADDING",(0,0),(-1,-1),4),
    ]))
    story.append(hdr)
    story.append(HRFlowable(width="100%", thickness=2, color=PRI_RL, spaceAfter=8))
    story.append(Paragraph("CERTIFICATE OF ANALYSIS", ST["ti"]))
    story.append(Paragraph(data.get("commonName",""), ST["pr"]))

    # Meta
    def mr(l,v):
        return [Paragraph(l, ST["ml"]), Paragraph(str(v or ""), ST["mv"])]

    ml, mr_list = [], []
    for (l1,v1),(l2,v2) in _meta_pairs(data):
        ml.append(mr(l1,v1)); mr_list.append(mr(l2,v2))

    meta_rows = [[l[0],l[1],r[0],r[1]] for l,r in zip(ml,mr_list)]
    mt = Table(meta_rows, colWidths=[cw*0.18,cw*0.32,cw*0.18,cw*0.32])
    mt.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),LIGHT_RL),
        ("GRID",(0,0),(-1,-1),0.5,BORDER_RL),
        ("TOPPADDING",(0,0),(-1,-1),4),
        ("BOTTOMPADDING",(0,0),(-1,-1),4),
        ("LEFTPADDING",(0,0),(-1,-1),6),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
    ]))
    story.append(mt)

    desc = data.get("description","")
    if desc:
        story.append(Spacer(1,4))
        story.append(Paragraph(f"<i>Description: {desc}</i>", ST["no"]))
    story.append(Spacer(1,8))

    # Params
    lo = _calc_layout(data)
    if lo["has_unit"]:
        ph  = [Paragraph("PARAMETER",ST["th"]), Paragraph("UNIT",ST["thc"]),
               Paragraph("MIN / MAX",ST["th"]),  Paragraph("RESULT",ST["thc"])]
        pcw = [cw*0.42, cw*0.10, cw*0.25, cw*0.23]
    elif lo["has_method"]:
        ph  = [Paragraph("PARAMETER",ST["th"]), Paragraph("SPECIFICATION",ST["th"]),
               Paragraph("RESULT",ST["thc"]),   Paragraph("METHOD",ST["thc"])]
        pcw = [cw*0.38, cw*0.27, cw*0.18, cw*0.17]
    else:
        ph  = [Paragraph("PARAMETER",ST["th"]), Paragraph("MIN / MAX",ST["th"]),
               Paragraph("RESULT",ST["thc"]),   Paragraph("",ST["thc"])]
        pcw = [cw*0.50, cw*0.26, cw*0.20, cw*0.04]

    p_rows = [ph]
    current_section = None
    for i, p in enumerate(data.get("parameters",[])):
        if lo["has_sections"] and p.get("section") and p["section"] != current_section:
            current_section = p["section"]
            p_rows.append([Paragraph(current_section.upper(), ST["se"]),"","",""])

        st   = p.get("status","pass")
        icon = "✓" if st=="pass" else "✗" if st=="fail" else "·"
        if lo["has_unit"]:
            row_d = [Paragraph(p.get("name",""),ST["pn"]),
                     Paragraph(p.get("unit",""),ST["pv"]),
                     Paragraph(p.get("min_max",""),ST["pv"]),
                     Paragraph(p.get("result",""),ST["pc"])]
        elif lo["has_method"]:
            row_d = [Paragraph(p.get("name",""),ST["pn"]),
                     Paragraph(p.get("min_max",""),ST["pv"]),
                     Paragraph(p.get("result",""),ST["pc"]),
                     Paragraph(p.get("method",""),ST["pv"])]
        else:
            row_d = [Paragraph(p.get("name",""),ST["pn"]),
                     Paragraph(p.get("min_max",""),ST["pv"]),
                     Paragraph(p.get("result",""),ST["pc"]),
                     Paragraph(icon,ST["pc"])]
        p_rows.append(row_d)

    pt = Table(p_rows, colWidths=pcw)
    pst = [
        ("BACKGROUND",(0,0),(-1,0),PRI_RL),
        ("TEXTCOLOR",(0,0),(-1,0),WHITE_RL),
        ("TOPPADDING",(0,0),(-1,0),7),("BOTTOMPADDING",(0,0),(-1,0),7),
        ("TOPPADDING",(0,1),(-1,-1),5),("BOTTOMPADDING",(0,1),(-1,-1),5),
        ("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),
        ("GRID",(0,0),(-1,-1),0.5,BORDER_RL),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
    ]
    for i in range(1,len(p_rows)):
        pst.append(("BACKGROUND",(0,i),(-1,i),ODD_RL if i%2==1 else WHITE_RL))
    pt.setStyle(TableStyle(pst))
    story.append(pt)
    story.append(Spacer(1,6))

    for txt in [data.get("notes",""),
                ("Storage: "+data.get("storageConditions","")).strip()]:
        if txt and txt not in ["Storage: ",""]:
            story.append(Paragraph(txt, ST["no"]))
            story.append(Spacer(1,4))

    story.append(HRFlowable(width="100%", thickness=1, color=BORDER_RL, spaceBefore=8))
    story.append(Spacer(1,6))
    today = date.today().strftime("%d.%m.%Y")
    fl = Table([
        [Paragraph("QA Manager", ST["qa"])],
        [Spacer(1,20)],
        [HRFlowable(width=55*mm, thickness=1, color=PRI_RL)],
        [Paragraph("Signature &amp; Date", ST["no"])],
    ], colWidths=[70*mm])
    fr = Table([[fl, Table([
        [Paragraph("ORIN Group s.r.o. — Internal QC Document", ST["ge"])],
        [Paragraph(f"Generated: {today}", ST["ge"])],
    ], colWidths=[cw-70*mm])]], colWidths=[70*mm, cw-70*mm])
    fr.setStyle(TableStyle([("VALIGN",(0,0),(-1,-1),"BOTTOM")]))
    story.append(fr)

    doc.build(story)
    with open(output_path,"wb") as f:
        f.write(buf.getvalue())
    # logo is loaded directly from base64, no temp file needed
    return output_path
